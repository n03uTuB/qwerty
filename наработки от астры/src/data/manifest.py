# -*- coding: utf-8 -*-
"""Формирование манифеста датасета и работа с ним.

Манифест: одна строка = одно уникальное изображение (одна анатомическая
область). Метки качества берутся из колонок «Итог» разметки, метки типов
нарушений — из отдельных критериев. Синтетические строки (артефакты v2)
помечаются ``synthetic=True`` и исключаются из всех метрик/порогов/отчёта.

Схема колонок сохранена из базового решения (``dxa_qc``): на неё опираются
:mod:`src.features`, :mod:`src.stack` и OOF-массивы.
"""
from __future__ import annotations

import hashlib
import os
import warnings
from typing import Dict, List

import numpy as np
import pandas as pd

from .. import config as C
from .. import dicom_io

warnings.filterwarnings("ignore")

# Индексы столбцов листа «Калибровка» (1-based)
_COL_STUDY = 2
_COL_SPINE = (3, 4, 5)          # укладка, ось, артефакты
_COL_FEMUR_R = (6, 7)           # позиционирование/ротация, ROI
_COL_FEMUR_L = (8, 9)
_COL_TOTAL = (10, 11, 12)       # итог: позвоночник, правое бедро, левое бедро


def _read_labels(xlsx_path: str) -> Dict[str, dict]:
    """Прочитать разметку из листа «Калибровка» (файл организатора)."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Калибровка"]
    labels: Dict[str, dict] = {}
    for r in range(3, ws.max_row + 1):
        uid = ws.cell(r, _COL_STUDY).value
        if uid is None:
            continue
        labels[str(uid)] = dict(
            spine=tuple(ws.cell(r, c).value for c in _COL_SPINE),
            femur_r=tuple(ws.cell(r, c).value for c in _COL_FEMUR_R),
            femur_l=tuple(ws.cell(r, c).value for c in _COL_FEMUR_L),
            totals=tuple(ws.cell(r, c).value for c in _COL_TOTAL),
        )
    return labels


def _quality_row(region: str, lab: dict):
    """Вернуть (quality, viol_vec[5], viol_mask[5]) для изображения области."""
    viol = np.full(C.N_VIOLATIONS, np.nan, dtype=np.float32)
    mask = np.zeros(C.N_VIOLATIONS, dtype=np.float32)
    quality = np.nan

    if region == C.REGION_SPINE:
        total = lab["totals"][0]
        crit = lab["spine"]
        names = C.REGION_CRITERIA[C.REGION_SPINE]
    else:
        idx = 1 if region == C.REGION_FEMUR_RIGHT else 2
        total = lab["totals"][idx]
        crit = lab["femur_r"] if region == C.REGION_FEMUR_RIGHT else lab["femur_l"]
        names = C.REGION_CRITERIA[region]

    if total is not None:
        quality = float(total)
    for name, val in zip(names, crit):
        if val is not None:
            viol[C.VIOLATION_IDX[name]] = float(val)
            mask[C.VIOLATION_IDX[name]] = 1.0
    return quality, viol, mask


def build_manifest(dataset_root: str, manifest_csv: str = C.MANIFEST_CSV,
                   cache_dir: str = C.CACHE_DIR, verbose: bool = True) -> pd.DataFrame:
    """Построить манифест по каталогу НД_для_обучения (или аналогу)."""
    from PIL import Image

    studies_root = os.path.join(dataset_root, "исследования")
    xlsx = os.path.join(dataset_root, "разметка.xlsx")
    if not os.path.isdir(studies_root):
        studies_root = dataset_root
    labels = _read_labels(xlsx) if os.path.isfile(xlsx) else {}

    os.makedirs(os.path.dirname(os.path.abspath(manifest_csv)), exist_ok=True)
    os.makedirs(cache_dir, exist_ok=True)

    rows: List[dict] = []
    n_mismatch = 0
    for study_uid, study_path in dicom_io.iter_studies(studies_root):
        files = dicom_io.find_dicom_files(study_path)
        if not files:
            continue
        uniq = dicom_io.dedupe_unique_images(files)
        regions = dicom_io.detect_region_study(uniq)
        copies = dicom_io.study_copies(files)
        lab = labels.get(study_uid)

        if lab is not None:
            n_groups = sum(
                1 for t in (lab["spine"], lab["femur_r"], lab["femur_l"])
                if any(v is not None for v in t)
            )
            if n_groups != len(uniq):
                n_mismatch += 1

        for (path, ds, arr), region in zip(uniq, regions):
            image_uid = str(getattr(ds, "SOPInstanceUID", "") or "")
            if not image_uid:
                image_uid = hashlib.md5(path.encode("utf-8")).hexdigest()

            quality, viol, mask = (np.nan,
                                   np.full(C.N_VIOLATIONS, np.nan, np.float32),
                                   np.zeros(C.N_VIOLATIONS, np.float32))
            if lab is not None:
                quality, viol, mask = _quality_row(region, lab)

            # Кэш изображения в PNG. to_display_uint8 применяет modality/VOI LUT —
            # тот же конвейер, что и при инференсе (нет рассинхрона train/infer).
            h = hashlib.md5(image_uid.encode("utf-8")).hexdigest()
            cache_path = os.path.join(cache_dir, h + ".png")
            if not os.path.isfile(cache_path):
                Image.fromarray(dicom_io.to_display_uint8(ds, arr)).save(cache_path)

            spacing_x, spacing_y = dicom_io.pixel_spacing(ds, arr)
            px_hash = hashlib.md5(np.ascontiguousarray(arr).tobytes()).hexdigest()
            row = dict(
                study_uid=study_uid,
                image_uid=image_uid,
                source_path=path,
                cache_path=cache_path,
                region=region,
                region_idx=C.REGION_TO_IDX[region],
                rows=int(arr.shape[0]),
                cols=int(arr.shape[1]),
                spacing_x=spacing_x,
                spacing_y=spacing_y,
                copies=int(copies.get(px_hash, 1)),
                quality=quality,
                synthetic=False,
            )
            for i, name in enumerate(C.VIOLATIONS):
                row["viol_" + name] = viol[i]
                row["mask_" + name] = mask[i]
            rows.append(row)

    df = pd.DataFrame(rows)
    save_manifest(df, manifest_csv)
    if verbose:
        print(f"[manifest] изображений: {len(df)}  исследований: {df.study_uid.nunique()}")
        print(f"[manifest] несовпадение числа областей: {n_mismatch}")
        print("[manifest] по областям:\n", df.region.value_counts().to_string())
        labeled = df.dropna(subset=["quality"])
        print("[manifest] размечено качеством:", len(labeled))
        print("[manifest] баланс quality:\n", labeled.quality.value_counts().to_string())
        for name in C.VIOLATIONS:
            sub = df.dropna(subset=["viol_" + name])
            if len(sub):
                print(f"    {name}: pos={int(sub['viol_'+name].sum())}/{len(sub)}")
    return df


def save_manifest(df: pd.DataFrame, manifest_csv: str = C.MANIFEST_CSV) -> None:
    """Сохранить манифест в CSV (UTF-8)."""
    os.makedirs(os.path.dirname(os.path.abspath(manifest_csv)), exist_ok=True)
    df.to_csv(manifest_csv, index=False, encoding="utf-8")


def load_manifest(manifest_csv: str = C.MANIFEST_CSV) -> pd.DataFrame:
    """Загрузить манифест и привести типы меток к числам."""
    df = pd.read_csv(manifest_csv)
    for c in ["quality"] + ["viol_" + n for n in C.VIOLATIONS]:
        if c in df:
            df[c] = pd.to_numeric(df[c], errors="coerce")
    if "synthetic" not in df:
        df["synthetic"] = False
    df["synthetic"] = df["synthetic"].fillna(False).astype(bool)
    return df


def real_mask(df: pd.DataFrame) -> np.ndarray:
    """Маска настоящих снимков: по ним считаются все метрики."""
    if "synthetic" not in df:
        return np.ones(len(df), dtype=bool)
    return ~df["synthetic"].fillna(False).values.astype(bool)