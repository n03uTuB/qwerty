"""Загрузка набора НД_для_обучения: изображения, метки эксперта, метаданные.

Особенности данных (проверено на 100 исследованиях, GE Lunar Prodigy Advance):
  * UID исследования — имя папки, теги обезличены (BodyPart, Laterality, PixelSpacing пусты);
  * одно изображение хранится в нескольких копиях; число копий — полезный признак;
  * порядок съёмки устойчив: позвоночник имеет InstanceNumber = 1 (91 из 95), дальше бёдра;
  * область определяется шириной кадра (300 — позвоночник, 280/248 — бедро) и порядком;
  * сторона бедра — по наклону оси кости внутри пары (меньший наклон — правое).
"""

from __future__ import annotations

import warnings
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pydicom

warnings.filterwarnings("ignore")

SPINE_MIN_WIDTH = 295
# Тега PixelSpacing нет, НО есть (0040,0303) Exposed Area — физический размер снятой области
# в мм. Он присутствует во всех 499 файлах набора и даёт размер пикселя для каждого снимка:
# медиана 0.607 × 0.608 мм, то есть пиксель квадратный (а не 0.6 × 1.05, как принято «на глаз»).
PIXEL_SPACING_MM = (0.6, 0.6)  # запасное значение, если тега нет
EXPOSED_AREA_TAG = (0x0040, 0x0303)

# Метки организатора (закрытый список)
LABEL_POSITIONING = "Некорректная укладка"
LABEL_AXIS = "Не выравнена ось позвоночника"
LABEL_FOREIGN = "Присутствуют посторонние предметы"
LABEL_ROI = "Некорректная область интереса"
LABELS = [LABEL_POSITIONING, LABEL_AXIS, LABEL_FOREIGN, LABEL_ROI]

# Колонки листа «Калибровка» (1-based)
_COL = dict(study=2, sp_pos=3, sp_axis=4, sp_art=5, fr_pos=6, fr_roi=7, fl_pos=8, fl_roi=9,
            total_spine=10, total_fr=11, total_fl=12)


@dataclass
class Image:
    study_uid: str
    image_uid: str
    path: Path
    array: np.ndarray
    region: str  # "spine" | "femur"
    side: str | None  # "l" | "r" для бедра
    instance: int
    copies: int  # сколько раз это же изображение лежит в исследовании
    n_images: int  # уникальных изображений в исследовании
    spacing: tuple[float, float] = PIXEL_SPACING_MM  # (мм по X, мм по Y) для этого снимка
    quality: int | None = None
    labels: dict[str, int] = field(default_factory=dict)  # метка организатора -> 0/1

    @property
    def key(self) -> str:
        return f"{self.study_uid}:{self.image_uid}"


def pixel_spacing(ds, array: np.ndarray) -> tuple[float, float]:
    """Размер пикселя из Exposed Area (0040,0303): [ширина, высота] снятой области в мм."""
    area = ds.get(EXPOSED_AREA_TAG)
    if area is None or len(area.value) < 2:
        return PIXEL_SPACING_MM
    width_mm, height_mm = float(area.value[0]), float(area.value[1])
    rows, cols = array.shape[:2]
    sx = width_mm / cols if cols and width_mm > 0 else PIXEL_SPACING_MM[0]
    sy = height_mm / rows if rows and height_mm > 0 else PIXEL_SPACING_MM[1]
    # защита от мусорных значений в отдельных файлах
    if not (0.2 <= sx <= 2.0 and 0.2 <= sy <= 2.0):
        return PIXEL_SPACING_MM
    return sx, sy


def read_labels(xlsx_path: str | Path) -> dict[str, dict]:
    import openpyxl

    ws = openpyxl.load_workbook(xlsx_path, data_only=True)["Калибровка"]
    labels = {}
    for row in range(3, ws.max_row + 1):
        uid = ws.cell(row, _COL["study"]).value
        if uid is None:
            continue
        values = {}
        for name, column in _COL.items():
            if name == "study":
                continue
            value = ws.cell(row, column).value
            values[name] = None if value is None else int(value)
        labels[str(uid)] = values
    return labels


def _bone_slope(arr: np.ndarray) -> float:
    """Наклон главной оси костной структуры: отличает левое бедро от правого."""
    values = arr.astype(np.float32)
    if values.max() <= values.min():
        return 0.0
    ys, xs = np.nonzero(values >= np.percentile(values, 92))
    if len(xs) < 10:
        return 0.0
    rows, cols = arr.shape
    coords = np.stack([xs / cols, ys / rows])
    eigvals, eigvecs = np.linalg.eigh(np.cov(coords))
    vx, vy = eigvecs[:, int(np.argmax(eigvals))]
    if vy < 0:
        vx, vy = -vx, -vy
    return float(vx / (vy + 1e-6))


def _assign_labels(image: Image, expert: dict) -> None:
    if image.region == "spine":
        image.quality = expert["total_spine"]
        image.labels = {LABEL_POSITIONING: expert["sp_pos"], LABEL_AXIS: expert["sp_axis"],
                        LABEL_FOREIGN: expert["sp_art"]}
    else:
        prefix = "fr" if image.side == "r" else "fl"
        image.quality = expert[f"total_{prefix}"]
        image.labels = {LABEL_POSITIONING: expert[f"{prefix}_pos"], LABEL_ROI: expert[f"{prefix}_roi"]}
    image.labels = {k: v for k, v in image.labels.items() if v is not None}


def load_dataset(root: str | Path, with_labels: bool = True) -> list[Image]:
    """Прочитать набор: одно изображение = одна анатомическая область."""
    root = Path(root)
    xlsx = next(root.glob("*.xlsx"), None)
    expert = read_labels(xlsx) if (with_labels and xlsx) else {}
    studies_root = next((p for p in root.iterdir() if p.is_dir()), root)

    raw: dict[str, list[tuple[int, str, np.ndarray]]] = defaultdict(list)
    for path in sorted(studies_root.rglob("*")):
        if not path.is_file() or path.suffix.lower() == ".xlsx":
            continue
        ds = pydicom.dcmread(path, force=True)
        try:
            array = np.asarray(ds.pixel_array)
        except Exception:
            continue
        study_uid = path.relative_to(studies_root).parts[0]
        image_uid = str(ds.get("SOPInstanceUID", "") or path.name)
        raw[study_uid].append((int(ds.get("InstanceNumber", 0) or 0), image_uid, array, path,
                              pixel_spacing(ds, array)))

    images: list[Image] = []
    for study_uid, items in raw.items():
        unique: dict[bytes, list] = defaultdict(list)
        for instance, image_uid, array, path, spacing in sorted(items, key=lambda t: t[0]):
            unique[array.tobytes()].append((instance, image_uid, array, path, spacing))
        groups = sorted(unique.values(), key=lambda g: g[0][0])

        study_images = []
        for group in groups:
            instance, image_uid, array, path, spacing = group[0]
            region = "spine" if array.shape[1] >= SPINE_MIN_WIDTH else "femur"
            study_images.append(Image(study_uid=study_uid, image_uid=image_uid, path=path,
                                      array=array, region=region, side=None, instance=instance,
                                      copies=len(group), n_images=len(groups), spacing=spacing))

        femurs = [im for im in study_images if im.region == "femur"]
        if len(femurs) == 2:  # в паре сторона определяется однозначно
            order = np.argsort([_bone_slope(im.array) for im in femurs])
            femurs[int(order[0])].side = "r"
            femurs[int(order[1])].side = "l"
        elif len(femurs) == 1:
            femurs[0].side = "l" if _bone_slope(femurs[0].array) >= 0 else "r"

        for image in study_images:
            if expert.get(study_uid):
                _assign_labels(image, expert[study_uid])
            images.append(image)
    return images


def summary(images: list[Image]) -> str:
    lines = [f"изображений: {len(images)}, исследований: {len({i.study_uid for i in images})}"]
    for region in ("spine", "femur"):
        subset = [i for i in images if i.region == region]
        labeled = [i for i in subset if i.quality is not None]
        positive = sum(1 for i in labeled if i.quality)
        lines.append(f"  {region}: {len(subset)} (с меткой {len(labeled)}, нарушений {positive})")
    for label in LABELS:
        subset = [i for i in images if label in i.labels]
        if subset:
            lines.append(f"  {label}: {sum(i.labels[label] for i in subset)} из {len(subset)}")
    return "\n".join(lines)
